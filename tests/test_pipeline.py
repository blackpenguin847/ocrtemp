"""렌더링 → 추출 → 출력 전체 경로를, Ollama 대신 가짜 클라이언트로 검증한다."""

import json

import pymupdf
import pytest
from openpyxl import load_workbook

from estimate_ocr import Extractor, export
from estimate_ocr.exporter import SHEET_COST, SHEET_RESULT, SHEET_REVIEW
from estimate_ocr.ollama_client import OllamaError
from estimate_ocr.pdf_render import collect_pdfs, render_pdf

GOOD_JSON = json.dumps(
    {
        "문서번호": "Q-2024-001",
        "원가": {
            "재료비": 3000,
            "노무비": 2000,
            "관리비": 500,
            "포장비": 300,
            "영업이익": 700,
            "합계": 6500,
        },
        "품목": [
            {"구분": "재료비", "품명": "레미콘", "규격": "25-24-150", "단위": "㎥", "수량": 2, "단가": 1500, "금액": 3000},
            {"구분": "노무비", "품명": "펌프카", "규격": "32m", "단위": "대", "수량": 1, "단가": 2000, "금액": 2000},
        ],
    },
    ensure_ascii=False,
)


class FakeClient:
    """`generate` / `generate_json` 만 흉내내는 스텁."""

    def __init__(self, responses, model="fake-vl"):
        self.responses = list(responses)
        self.model = model
        self.calls = []

    def _next(self, kind):
        self.calls.append(kind)
        if not self.responses:
            raise AssertionError("예상보다 많은 호출이 발생했습니다")
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    def generate(self, prompt, *, images=None, system=None, json_mode=False):
        return self._next("generate")

    def generate_json(self, prompt, *, images=None, system=None):
        from estimate_ocr.ollama_client import loads_lenient

        return loads_lenient(self._next("generate_json"))


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / "견적서.pdf"
    document = pymupdf.open()
    for index in range(2):
        page = document.new_page()
        page.insert_text((72, 100), f"ESTIMATE SHEET p{index + 1}", fontsize=18)
        page.insert_text((72, 140), "item | qty | price | amount", fontsize=11)
    document.save(path)
    document.close()
    return path


def test_collect_pdfs_finds_files_and_folders(sample_pdf, tmp_path):
    assert collect_pdfs(sample_pdf) == [sample_pdf]
    assert collect_pdfs(tmp_path) == [sample_pdf]
    with pytest.raises(FileNotFoundError):
        collect_pdfs(tmp_path / "없음")


def test_render_pdf_respects_pages_and_max_edge(sample_pdf):
    pages = list(render_pdf(sample_pdf, dpi=100, max_edge=400, pages=[2]))
    assert [page.page for page in pages] == [2]
    assert max(pages[0].image.size) <= 400
    assert pages[0].image.mode == "L"


def test_extract_pdf_happy_path(sample_pdf):
    client = FakeClient([GOOD_JSON, GOOD_JSON])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf)

    assert len(docs) == 2
    assert [doc.page for doc in docs] == [1, 2]
    assert len(docs[0].items) == 2
    assert docs[0].items[0].amount == 3000
    assert docs[0].items[0].category == "재료비"
    assert docs[0].costs.material == 3000
    assert docs[0].costs.profit == 700
    assert docs[0].cost_total() == 6500
    assert docs[0].validate() == []
    assert docs[0].attempts == 1
    assert docs[0].model == "fake-vl"


def test_extract_falls_back_to_simple_prompt(sample_pdf):
    client = FakeClient(["설명만 하고 JSON이 없습니다", GOOD_JSON], model="fake-vl")
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])

    assert docs[0].attempts == 2
    assert len(docs[0].items) == 2


def test_extract_falls_back_to_text_then_json(sample_pdf):
    client = FakeClient(
        [
            "JSON 아님",                       # 1차 (json)
            "여전히 JSON 아님",                 # 2차 (json)
            "레미콘 | 2 | 1,500 | 3,000",      # 3차 원문
            GOOD_JSON,                         # 3차 구조화
        ]
    )
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])

    assert docs[0].attempts == 3
    assert len(docs[0].items) == 2
    assert "레미콘" in docs[0].raw_text


def test_extract_records_error_when_all_attempts_fail(sample_pdf):
    client = FakeClient(["없음", "없음", "", ])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])

    assert docs[0].error
    assert docs[0].items == []
    assert any(issue.code == "parse_error" for issue in docs[0].validate())


def test_ollama_error_is_recorded_not_raised(sample_pdf):
    client = FakeClient([OllamaError("연결 실패")] * 3)
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])
    assert "연결 실패" in docs[0].error


def test_raw_text_mode(sample_pdf, tmp_path):
    from estimate_ocr.exporter import write_raw_text

    client = FakeClient(["레미콘 | 2 | 1,500 | 3,000"])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1], raw_text=True)
    assert docs[0].raw_text.startswith("레미콘")
    assert docs[0].items == []

    written = write_raw_text(docs, tmp_path / "out")
    assert written[0].read_text(encoding="utf-8").startswith("레미콘")


def test_debug_images_are_saved(sample_pdf, tmp_path):
    debug_dir = tmp_path / "out" / "_debug"
    client = FakeClient([GOOD_JSON])
    Extractor(client, debug_dir=debug_dir, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])
    assert (debug_dir / "견적서_p001.png").exists()


def test_export_all_formats(sample_pdf, tmp_path):
    bad_json = json.dumps(
        {"품목": [{"품명": "", "수량": 2, "단가": 1000, "금액": 5}]}, ensure_ascii=False
    )
    client = FakeClient([GOOD_JSON, bad_json])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf)

    written = export(docs, tmp_path / "out", "all")
    names = {path.name for path in written}
    assert {"result.json", "result.csv", "result.xlsx"} <= names

    payload = json.loads((tmp_path / "out" / "result.json").read_text(encoding="utf-8"))
    assert payload["문서수"] == 2
    assert payload["품목수"] == 3
    assert payload["문서"][0]["원가"]["영업이익"] == 700

    workbook = load_workbook(tmp_path / "out" / "result.xlsx")
    sheet = workbook[SHEET_RESULT]
    assert sheet.max_row == 4  # 헤더 + 품목 3건
    assert [cell.value for cell in sheet[1]][:5] == [
        "원본PDF", "페이지", "문서번호", "구분", "품명",
    ]

    # 마지막 행(품명 없음 + 금액 불일치)은 노란색으로 표시된다
    flagged = sheet.cell(row=4, column=1)
    assert flagged.fill.fgColor.rgb.endswith("FFF2CC")
    # 정상 행은 채우기 없음
    assert not sheet.cell(row=2, column=1).fill.fgColor.rgb.endswith("FFF2CC")

    review = workbook[SHEET_REVIEW]
    assert review.max_row >= 3
    assert {"empty_name", "amount_mismatch"} <= {row[4] for row in review.iter_rows(min_row=2, values_only=True)}


def test_export_marks_failed_documents(sample_pdf, tmp_path):
    client = FakeClient(["없음", "없음", ""])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])
    export(docs, tmp_path / "out", "xlsx")

    sheet = load_workbook(tmp_path / "out" / "result.xlsx")[SHEET_RESULT]
    assert sheet.max_row == 2
    assert sheet.cell(row=2, column=1).fill.fgColor.rgb.endswith("F8CBAD")


# --------------------------------------------------------------------------
# 비목(재료비/노무비/관리비/포장비/영업이익)
# --------------------------------------------------------------------------
COST_ONLY_JSON = json.dumps(
    {
        "문서번호": "C-2024-007",
        "품목": [
            {"품명": "재료비", "금액": "1,000,000"},
            {"품명": "노무비", "금액": "400,000"},
            {"품명": "일반관리비", "금액": "70,000"},
            {"품명": "포장비", "금액": "30,000"},
            {"품명": "이윤", "금액": "150,000"},
            {"품명": "합계", "금액": "1,650,000"},
        ],
    },
    ensure_ascii=False,
)


def test_cost_summary_sheet_is_written(sample_pdf, tmp_path):
    client = FakeClient([GOOD_JSON])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])
    export(docs, tmp_path / "out", "all")

    workbook = load_workbook(tmp_path / "out" / "result.xlsx")
    assert workbook.sheetnames == ["추출결과", "원가요약", "검수"]

    sheet = workbook[SHEET_COST]
    headers = [cell.value for cell in sheet[1]]
    for category in ("재료비", "노무비", "관리비", "포장비", "영업이익"):
        assert category in headers
    row = dict(zip(headers, [cell.value for cell in sheet[2]]))
    assert row["재료비"] == 3000
    assert row["영업이익"] == 700
    assert row["합계"] == 6500
    assert row["비목합계"] == 6500

    assert (tmp_path / "out" / "result_원가요약.csv").exists()


def test_category_summary_rows_move_out_of_item_table(sample_pdf, tmp_path):
    """'재료비 | 1,000,000' 같은 집계 행은 품목이 아니라 원가 요약으로 간다."""
    client = FakeClient([COST_ONLY_JSON])
    docs = Extractor(client, ocr_mode="off").extract_pdf(sample_pdf, pages=[1])
    doc = docs[0]

    assert doc.items == []  # 집계 행만 있는 문서
    assert doc.costs.material == 1_000_000
    assert doc.costs.overhead == 70_000      # 일반관리비 → 관리비
    assert doc.costs.profit == 150_000       # 이윤 → 영업이익
    assert doc.costs.total == 1_650_000
    assert doc.validate() == []              # 합계가 맞으므로 경고 없음

    # 세부 품목이 없어도 추출 실패(주황색)로 표시하지 않는다
    export(docs, tmp_path / "out", "xlsx")
    sheet = load_workbook(tmp_path / "out" / "result.xlsx")[SHEET_RESULT]
    assert not sheet.cell(row=2, column=1).fill.fgColor.rgb.endswith("F8CBAD")
