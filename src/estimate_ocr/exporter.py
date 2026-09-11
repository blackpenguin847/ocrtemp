"""JSON / CSV / XLSX 출력.

xlsx 는 두 개의 시트를 만든다.

- **추출결과**: 한 행 = 명세서의 한 품목. 경고가 걸린 행은 노란색으로 표시한다.
- **검수**: 자동 검증에서 걸린 경고 목록.
"""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Document

log = logging.getLogger(__name__)

FORMATS = ("json", "csv", "xlsx", "all")

RESULT_HEADERS = [
    "원본PDF", "페이지", "문서번호", "품명", "규격", "단위",
    "수량", "단가", "금액", "비고",
]
REVIEW_HEADERS = ["원본PDF", "페이지", "행", "품명", "코드", "내용"]

SHEET_RESULT = "추출결과"
SHEET_REVIEW = "검수"

WARN_FILL = PatternFill("solid", fgColor="FFF2CC")   # 경고 행 (노란색)
ERROR_FILL = PatternFill("solid", fgColor="F8CBAD")  # 추출 실패 행 (주황색)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
NUMBER_FORMAT = "#,##0.###"


def result_rows(documents: Sequence[Document]) -> list[dict[str, Any]]:
    """한 행 = 한 품목. 품목이 없는 문서도 흔적을 남긴다."""
    rows: list[dict[str, Any]] = []
    for document in documents:
        if not document.items:
            rows.append(
                {
                    "원본PDF": document.source,
                    "페이지": document.page,
                    "문서번호": document.doc_no,
                    "품명": "",
                    "규격": "",
                    "단위": "",
                    "수량": None,
                    "단가": None,
                    "금액": None,
                    "비고": document.error or "품목 없음",
                    "_issues": [issue.message for issue in document.validate()],
                    "_failed": True,
                }
            )
            continue

        issues_by_row: dict[int, list[str]] = {}
        for issue in document.validate():
            if issue.row:
                issues_by_row.setdefault(issue.row, []).append(issue.message)

        for index, item in enumerate(document.items, start=1):
            rows.append(
                {
                    "원본PDF": document.source,
                    "페이지": document.page,
                    "문서번호": document.doc_no,
                    "품명": item.name,
                    "규격": item.spec,
                    "단위": item.unit,
                    "수량": item.qty,
                    "단가": item.unit_price,
                    "금액": item.amount,
                    "비고": item.note,
                    "_issues": issues_by_row.get(index, []),
                    "_failed": False,
                }
            )
    return rows


def review_rows(documents: Sequence[Document]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for document in documents:
        for issue in document.validate():
            name = ""
            if issue.row and 1 <= issue.row <= len(document.items):
                name = document.items[issue.row - 1].name
            rows.append(
                {
                    "원본PDF": document.source,
                    "페이지": document.page,
                    "행": issue.row or "",
                    "품명": name,
                    "코드": issue.code,
                    "내용": issue.message,
                }
            )
    return rows


# --------------------------------------------------------------------------
# 개별 포맷
# --------------------------------------------------------------------------
def write_json(documents: Sequence[Document], path: Path) -> Path:
    payload = {
        "문서수": len(documents),
        "품목수": sum(len(document.items) for document in documents),
        "경고수": sum(len(document.validate()) for document in documents),
        "문서": [document.to_dict() for document in documents],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def write_csv(documents: Sequence[Document], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = result_rows(documents)
    # 엑셀에서 바로 열어도 한글이 깨지지 않도록 BOM 포함
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=RESULT_HEADERS + ["검수"])
        writer.writeheader()
        for row in rows:
            record = {key: row.get(key, "") for key in RESULT_HEADERS}
            record["검수"] = " / ".join(row["_issues"])
            writer.writerow(record)
    return path


def write_review_csv(documents: Sequence[Document], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REVIEW_HEADERS)
        writer.writeheader()
        writer.writerows(review_rows(documents))
    return path


def write_xlsx(documents: Sequence[Document], path: Path) -> Path:
    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = SHEET_RESULT
    _write_sheet(result_sheet, RESULT_HEADERS)

    numeric_columns = {RESULT_HEADERS.index(name) + 1 for name in ("수량", "단가", "금액")}

    for row in result_rows(documents):
        values = [row.get(header) for header in RESULT_HEADERS]
        result_sheet.append(values)
        index = result_sheet.max_row

        for column in numeric_columns:
            result_sheet.cell(row=index, column=column).number_format = NUMBER_FORMAT

        if row["_issues"] or row["_failed"]:
            fill = ERROR_FILL if row["_failed"] else WARN_FILL
            for column in range(1, len(RESULT_HEADERS) + 1):
                result_sheet.cell(row=index, column=column).fill = fill
            if row["_issues"]:
                # 값은 원본 그대로 두고, 사유는 셀 메모로만 붙인다 (검수 시트에도 동일 내용)
                target = result_sheet.cell(row=index, column=RESULT_HEADERS.index("품명") + 1)
                target.comment = Comment(" / ".join(row["_issues"]), "estimate-ocr")

    review_sheet = workbook.create_sheet(SHEET_REVIEW)
    _write_sheet(review_sheet, REVIEW_HEADERS)
    for row in review_rows(documents):
        review_sheet.append([row[header] for header in REVIEW_HEADERS])
    if review_sheet.max_row == 1:
        review_sheet.append(["", "", "", "", "", "경고 없음"])

    for sheet in (result_sheet, review_sheet):
        _autosize(sheet)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _write_sheet(sheet: Any, headers: Sequence[str]) -> None:
    sheet.append(list(headers))
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(sheet: Any, *, minimum: int = 8, maximum: int = 48) -> None:
    for column_cells in sheet.columns:
        longest = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(minimum, longest + 2), maximum)


def write_raw_text(documents: Sequence[Document], directory: Path) -> list[Path]:
    """--raw-text 모드 결과를 페이지별 .txt 로 저장한다."""
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for document in documents:
        stem = Path(document.source).stem or "document"
        path = directory / f"{stem}_p{document.page:03d}.txt"
        path.write_text(document.raw_text or document.error, encoding="utf-8")
        written.append(path)
    return written


# --------------------------------------------------------------------------
# 진입점
# --------------------------------------------------------------------------
def export(
    documents: Sequence[Document],
    output_dir: str | Path,
    fmt: str = "xlsx",
    *,
    stem: str = "result",
) -> list[Path]:
    """`fmt` 에 해당하는 파일을 모두 쓰고 생성된 경로를 반환한다."""
    if fmt not in FORMATS:
        raise ValueError(f"알 수 없는 출력 형식: {fmt} (가능: {', '.join(FORMATS)})")

    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    wanted = ("json", "csv", "xlsx") if fmt == "all" else (fmt,)

    written: list[Path] = []
    if "json" in wanted:
        written.append(write_json(documents, directory / f"{stem}.json"))
    if "csv" in wanted:
        written.append(write_csv(documents, directory / f"{stem}.csv"))
        written.append(write_review_csv(documents, directory / f"{stem}_검수.csv"))
    if "xlsx" in wanted:
        written.append(write_xlsx(documents, directory / f"{stem}.xlsx"))

    for path in written:
        log.info("저장: %s", path)
    return written
